"""Build the NETRA solution-presentation deck from docs/presentation-outline.md.

Usage:
    python tools/build_deck.py            # build docs/submission/presentation.pptx
    python tools/build_deck.py --check     # build, then run overflow/slide-count assertions

Parser contract (see task-C2 brief):
    - Slides are ``## Slide N — Title`` headings.
    - ``- `` lines under a slide are bullets.
    - A line starting ``img: name.png`` inserts that screenshot, right-aligned,
      fit within a 6.2in x 4.4in box (missing images are skipped gracefully).
    - ``Say:`` lines (or a paragraph starting with ``Say:``) become speaker notes.
    - Everything else (prose, tables, blockquotes) is ignored by the deck build;
      it stays in the outline as authoring context.

No new runtime dependency is added to the product; python-pptx and Pillow are
docs/test-only tooling per global-constraints.md.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent.parent
OUTLINE_PATH = ROOT / "docs" / "presentation-outline.md"
SCREENSHOTS_DIR = ROOT / "docs" / "submission" / "screenshots"
OUTPUT_PATH = ROOT / "docs" / "submission" / "presentation.pptx"

BG = RGBColor(0x07, 0x0D, 0x1A)
TEXT = RGBColor(0xE6, 0xEC, 0xF7)
ACCENT = RGBColor(0xFF, 0x6B, 0x00)
MUTED = RGBColor(0x8F, 0xA1, 0xC2)

TITLE_PT = 36
BODY_PT = 20
MAX_BULLETS = 6
MAX_BULLET_CHARS = 120

SLIDE_HEADING_RE = re.compile(r"^##\s*Slide\s+(\d+)\s*[—-]\s*(.+?)\s*$")
IMG_RE = re.compile(r"^img:\s*(\S+)\s*$", re.IGNORECASE)
SAY_RE = re.compile(r"^\*?Speaker note:\s*(.+?)\*?$", re.IGNORECASE)
SAY_PREFIX_RE = re.compile(r"^Say:\s*(.+)$", re.IGNORECASE)


def _pick_font() -> str:
    """Return 'Plus Jakarta Sans' if a font of that name looks installed, else Calibri."""
    try:
        import matplotlib.font_manager as fm  # optional; not guaranteed installed

        names = {f.name for f in fm.fontManager.ttflist}
        if "Plus Jakarta Sans" in names:
            return "Plus Jakarta Sans"
    except Exception:
        pass
    # Fallback: check common Windows font-file locations.
    candidates = [
        Path("C:/Windows/Fonts/PlusJakartaSans-Regular.ttf"),
        Path.home() / "AppData/Local/Microsoft/Windows/Fonts/PlusJakartaSans-Regular.ttf",
    ]
    if any(c.exists() for c in candidates):
        return "Plus Jakarta Sans"
    return "Calibri"


FONT = _pick_font()


class Slide:
    def __init__(self, number: int, title: str):
        self.number = number
        self.title = title
        self.bullets: list[str] = []
        self.img: str | None = None
        self.notes: list[str] = []


def parse_outline(text: str) -> list[Slide]:
    lines = text.splitlines()
    slides: list[Slide] = []
    current: Slide | None = None
    in_appendix = False

    for raw in lines:
        line = raw.rstrip("\n")
        stripped = line.strip()

        if stripped.startswith("## Appendix"):
            if current is not None:
                slides.append(current)
                current = None
            in_appendix = True
            continue
        if in_appendix:
            continue

        m = SLIDE_HEADING_RE.match(stripped)
        if m:
            if current is not None:
                slides.append(current)
            current = Slide(int(m.group(1)), m.group(2))
            continue

        if current is None:
            continue

        img_m = IMG_RE.match(stripped)
        if img_m:
            current.img = img_m.group(1)
            continue

        say_m = SAY_PREFIX_RE.match(stripped) or SAY_RE.match(stripped)
        if say_m:
            current.notes.append(say_m.group(1).strip())
            continue

        if stripped.startswith("- "):
            bullet = stripped[2:].strip()
            bullet = re.sub(r"\*\*(.+?)\*\*", r"\1", bullet)  # strip bold markers
            current.bullets.append(bullet)
            continue

        # everything else (tables, prose, blank lines) ignored per contract

    if current is not None:
        slides.append(current)

    return slides


def _set_background(slide, prs):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _add_title(slide, prs, text: str):
    left, top, width, height = Inches(0.5), Inches(0.35), Inches(9.0), Inches(1.0)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(TITLE_PT)
    run.font.bold = True
    run.font.name = FONT
    run.font.color.rgb = ACCENT
    return box


def _add_bullets(slide, bullets: list[str], has_image: bool):
    width = Inches(6.0) if has_image else Inches(11.7)
    left, top, height = Inches(0.5), Inches(1.5), Inches(5.5)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"\u2022 {bullet}"
        run.font.size = Pt(BODY_PT)
        run.font.name = FONT
        run.font.color.rgb = TEXT
        p.space_after = Pt(10)
    return box


def _add_image(slide, prs, image_name: str) -> bool:
    path = SCREENSHOTS_DIR / image_name
    if not path.exists():
        return False
    from PIL import Image

    with Image.open(path) as im:
        w, h = im.size
    box_w, box_h = Inches(6.2), Inches(4.4)
    aspect = w / h
    box_aspect = box_w / box_h
    if aspect > box_aspect:
        draw_w = box_w
        draw_h = Emu(int(box_w / aspect))
    else:
        draw_h = box_h
        draw_w = Emu(int(box_h * aspect))

    slide_w = prs.slide_width
    right_margin = Inches(0.5)
    left = slide_w - right_margin - draw_w
    top = Inches(1.5) + (box_h - draw_h) // 2
    slide.shapes.add_picture(str(path), left, top, width=draw_w, height=draw_h)
    return True


def build(check: bool = False) -> Presentation:
    text = OUTLINE_PATH.read_text(encoding="utf-8")
    slides = parse_outline(text)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    image_status: list[tuple[int, str, str]] = []  # (number, name, status)

    for s in slides:
        slide = prs.slides.add_slide(blank_layout)
        _set_background(slide, prs)
        _add_title(slide, prs, f"{s.number}. {s.title}")

        has_image = bool(s.img)
        real_image = False
        if s.img:
            real_image = _add_image(slide, prs, s.img)
            status = "real" if real_image else "MISSING (skipped)"
            image_status.append((s.number, s.img, status))
            has_image = real_image

        if s.bullets:
            _add_bullets(slide, s.bullets, has_image)

        if s.notes:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_tf.text = "\n".join(s.notes)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))

    print(f"Wrote {OUTPUT_PATH} with {len(slides)} slides")
    print("Screenshot status:")
    for number, name, status in image_status:
        print(f"  slide {number}: {name} -> {status}")

    if check:
        run_checks(slides, image_status)

    return prs


def run_checks(slides: list[Slide], image_status: list[tuple[int, str, str]]):
    errors = []

    if len(slides) != 17:
        errors.append(f"expected 17 slides, found {len(slides)}")

    for s in slides:
        if len(s.bullets) > MAX_BULLETS:
            errors.append(f"slide {s.number} ({s.title}) has {len(s.bullets)} bullets (> {MAX_BULLETS})")
        for b in s.bullets:
            if len(b) > MAX_BULLET_CHARS:
                errors.append(
                    f"slide {s.number} ({s.title}) bullet exceeds {MAX_BULLET_CHARS} chars: {b[:60]}..."
                )

    if errors:
        print("\nCHECK FAILED:")
        for e in errors:
            print(f"  - {e}")
        sys.exit(1)

    print(f"\nCHECK OK: {len(slides)} slides, no overflow heuristics tripped.")
    real = sum(1 for _, _, status in image_status if status == "real")
    missing = sum(1 for _, _, status in image_status if status != "real")
    print(f"Screenshots: {real} real, {missing} placeholder/missing.")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true", help="run slide-count and overflow assertions")
    args = parser.parse_args()
    build(check=args.check)


if __name__ == "__main__":
    main()
